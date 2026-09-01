"""Structural analysis of a workbook, before anything knows what the data means.

Everything here is deterministic. The probe answers four questions per sheet:
which rows are section banners, which columns hold data, which row is the
header, and how the sheet is laid out. Getting these wrong silently corrupts
every number downstream, so each decision is scored rather than assumed.

The central mechanic is keeping the *expanded* grid and the *originally
populated* mask side by side. Merged ranges have to be expanded so a module
label reaches all its rows, but a cell that exists only because of that
expansion is not evidence that a row contains data.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any, Optional

import openpyxl

from .models import SheetProbe, SheetShape

# A header cell is a short label. Anything longer is prose, which means we are
# looking at a data row that happens to start with text.
MAX_HEADER_LEN = 60

# Columns thinner than this share of the densest column are annotations parked
# outside the main block, not fields. Relative rather than absolute, so a
# legitimately sparse field (a bug-link column used four times) survives while
# a stray counter block does not.
MIN_RELATIVE_DENSITY = 0.05

HEADER_SEARCH_DEPTH = 25


@dataclass
class Sheet:
    """A sheet as two parallel grids: values, and where they really came from."""

    name: str
    values: list[list[Any]]
    original: list[list[bool]]
    banner_rows: set[int] = field(default_factory=set)

    @property
    def height(self) -> int:
        return len(self.values)

    def row_originals(self, r: int, cols: list[int]) -> list[int]:
        """Data columns of row `r` populated in the source, not by expansion."""
        if r >= len(self.values):
            return []
        return [
            c for c in cols
            if c < len(self.values[r])
            and self.original[r][c]
            and not _is_blank(self.values[r][c])
        ]


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _looks_like_label(v: Any) -> bool:
    return (
        isinstance(v, str)
        and bool(v.strip())
        and len(v.strip()) <= MAX_HEADER_LEN
        and not v.strip().replace(".", "").isdigit()
    )


def read_sheets(path: str) -> list[Sheet]:
    """Load every sheet, expanding merged ranges but recording what was real."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    sheets: list[Sheet] = []

    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            sheets.append(Sheet(ws.title, [], []))
            continue

        width = max(len(r) for r in rows)
        for r in rows:
            r.extend([None] * (width - len(r)))

        original = [
            [not _is_blank(v) for v in row] for row in rows
        ]

        for rng in ws.merged_cells.ranges:
            r0, c0 = rng.min_row - 1, rng.min_col - 1
            if r0 >= len(rows) or c0 >= width:
                continue
            value = rows[r0][c0]
            if value is None:
                continue
            for r in range(rng.min_row - 1, min(rng.max_row, len(rows))):
                for c in range(rng.min_col - 1, min(rng.max_col, width)):
                    if r == r0 and c == c0:
                        continue
                    rows[r][c] = value

        sheets.append(Sheet(ws.title, rows, original))

    wb.close()
    return sheets


def find_data_columns(sheet: Sheet) -> list[int]:
    """Columns belonging to the sheet's main block.

    Measured on the expanded grid so a merged module column counts as populated
    throughout its span, then filtered relative to the densest column.
    """
    if not sheet.values:
        return []

    width = max((len(r) for r in sheet.values), default=0)
    counts = [0] * width
    for row in sheet.values:
        for c, v in enumerate(row):
            if not _is_blank(v):
                counts[c] += 1

    if not counts or max(counts) == 0:
        return []

    threshold = max(2, int(max(counts) * MIN_RELATIVE_DENSITY))
    return [c for c, n in enumerate(counts) if n >= threshold]


def find_banner_rows(sheet: Sheet, data_cols: list[int]) -> set[int]:
    """Rows that label a section rather than carry a record.

    Both spellings appear in the wild and both reduce to the same signal: the
    source populated exactly one cell, at the left edge of the block. Whether
    the author then merged it across the full width is irrelevant.
    """
    if not data_cols:
        return set()

    leftmost = data_cols[0]
    banners: set[int] = set()
    for r in range(sheet.height):
        originals = sheet.row_originals(r, data_cols)
        if len(originals) == 1 and originals[0] == leftmost:
            banners.add(r)
    return banners


def detect_header_row(sheet: Sheet, data_cols: list[int]) -> Optional[int]:
    """Pick the row that best behaves like a header.

    Scored on how many data columns it labels, whether those labels are
    distinct, and whether the rows beneath look different from it. Rows that
    merely echo the candidate — the lower half of a two-row merged header — are
    skipped when measuring that contrast, or a merged header would score itself
    down and hand the decision to its own duplicate.
    """
    if not sheet.values or not data_cols:
        return None

    best_row, best_score = None, 0.0

    for i in range(min(HEADER_SEARCH_DEPTH, sheet.height)):
        row = sheet.values[i]
        labels = [
            row[c] for c in data_cols
            if c < len(row) and _looks_like_label(row[c])
        ]
        if len(labels) < 2:
            continue

        header_values = {
            str(row[c]).strip().lower()
            for c in data_cols
            if c < len(row) and not _is_blank(row[c])
        }
        coverage = len(labels) / len(data_cols)
        distinct = len({str(v).strip().lower() for v in labels}) / len(labels)

        ratios = []
        for j in range(i + 1, min(i + 10, sheet.height)):
            cells = [
                sheet.values[j][c] for c in data_cols
                if c < len(sheet.values[j]) and not _is_blank(sheet.values[j][c])
            ]
            if not cells:
                continue
            below = {str(v).strip().lower() for v in cells}
            if below <= header_values:      # an echo of this header, not data
                continue
            ratios.append(
                sum(1 for v in cells if _looks_like_label(v)) / len(cells)
            )
            if len(ratios) >= 5:
                break
        contrast = 1.0 - (sum(ratios) / len(ratios) if ratios else 0.0)

        score = coverage * 2.0 + distinct + contrast
        if score > best_score:
            best_row, best_score = i, score

    return best_row


def body_rows(sheet: Sheet, data_cols: list[int], header_row: int) -> list[int]:
    """Row indices holding actual records.

    A row qualifies on originally-populated cells only. Rows carrying nothing
    but the tail of a merged column — the 16 empty rows below the last test
    case, still holding their module label — are not records.
    """
    return [
        r for r in range(header_row + 1, sheet.height)
        if r not in sheet.banner_rows and sheet.row_originals(r, data_cols)
    ]


def _value_domain(
    sheet: Sheet, col: int, rows: list[int], limit: int = 400
) -> set[str]:
    seen: set[str] = set()
    for r in rows[:limit]:
        row = sheet.values[r]
        if col >= len(row):
            continue
        v = row[col]
        if _is_blank(v) or isinstance(v, (int, float, date, datetime)):
            continue
        seen.add(str(v).strip().lower())
    return seen


def classify_shape(
    sheet: Sheet, data_cols: list[int], header_row: Optional[int], rows: list[int]
) -> SheetShape:
    """Records, a matrix, a summary grid, or noise.

    The matrix test is a shared vocabulary: if most non-label columns draw their
    values from one small common set, the columns are dimensions and the cells
    are results. Record sheets fail this because their fields are unrelated —
    severity values and status values have nothing in common.
    """
    if header_row is None or len(data_cols) < 2 or len(rows) < 3:
        return SheetShape.NOISE

    candidates = data_cols[1:]
    if len(candidates) >= 3:
        domains = {c: _value_domain(sheet, c, rows) for c in candidates}
        enumerable = {c: d for c, d in domains.items() if d and len(d) <= 8}
        if len(enumerable) >= 3:
            # The shared vocabulary is what recurs *across* columns, not the
            # union of everything. A sparsely-filled comment column can hold
            # only six distinct strings and look enumerable; pooling by union
            # would let it drag its prose into the vocabulary and sink the test.
            spread: dict[str, int] = {}
            for d in enumerable.values():
                for v in d:
                    spread[v] = spread.get(v, 0) + 1
            shared = {
                v for v, n in sorted(spread.items(), key=lambda kv: -kv[1])[:8]
                if n / len(enumerable) >= 0.3
            }
            if shared:
                agreeing = sum(
                    1 for d in enumerable.values()
                    if len(d & shared) / len(d) >= 0.8
                )
                if agreeing / len(candidates) >= 0.6:
                    return SheetShape.MATRIX

    return SheetShape.GRID if len(rows) < 20 else SheetShape.RECORD


def probe_sheet(sheet: Sheet) -> SheetProbe:
    # First pass locates the header using every dense column, including any
    # summary block parked beside the table.
    data_cols = find_data_columns(sheet)
    sheet.banner_rows = find_banner_rows(sheet, data_cols)
    header_row = detect_header_row(sheet, data_cols)

    if header_row is not None:
        # Second pass keeps only columns that carry records *beneath* the
        # header. A counter block sitting above it is structure, not a field,
        # and dropping it here also sharpens the header scoring on re-run.
        rows = body_rows(sheet, data_cols, header_row)
        populated = {
            c for r in rows for c in sheet.row_originals(r, data_cols)
        }
        if populated and len(populated) < len(data_cols):
            data_cols = [c for c in data_cols if c in populated]
            sheet.banner_rows = find_banner_rows(sheet, data_cols)
            header_row = detect_header_row(sheet, data_cols)

    rows = body_rows(sheet, data_cols, header_row) if header_row is not None else []
    shape = classify_shape(sheet, data_cols, header_row, rows)

    return SheetProbe(
        name=sheet.name,
        shape=shape,
        header_row=(header_row + 1) if header_row is not None else None,
        data_columns=data_cols,
        first_data_row=(rows[0] + 1) if rows else None,
        last_data_row=(rows[-1] + 1) if rows else None,
        row_count=len(rows),
    )


def fingerprint(sheets: list[Sheet], probes: list[SheetProbe]) -> str:
    """Stable hash of the workbook's structure, ignoring its contents.

    Keys the mapping cache, so the same layout with different data resolves its
    columns without any model involvement.
    """
    by_name = {s.name: s for s in sheets}
    parts: list[str] = []
    for p in probes:
        parts.append(p.name)
        if p.header_row is None:
            continue
        row = by_name[p.name].values[p.header_row - 1]
        parts.extend(
            str(row[c]).strip().lower()
            if c < len(row) and row[c] is not None else ""
            for c in p.data_columns
        )
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:16]


def probe_workbook(path: str) -> tuple[list[Sheet], list[SheetProbe], str]:
    sheets = read_sheets(path)
    probes = [probe_sheet(s) for s in sheets]
    return sheets, probes, fingerprint(sheets, probes)
